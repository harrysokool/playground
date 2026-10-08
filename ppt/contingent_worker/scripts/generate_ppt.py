"""Generate the GIA HK 2027 Contingent Worker business case deck from the AIA template.

Usage (from ppt/contingency_worker/):
    python3 scripts/generate_ppt.py                     # writes current/<DECK_NAME>.pptx
    python3 scripts/generate_ppt.py --output out.pptx   # write elsewhere, e.g. to check reproduction

An existing output file is never overwritten unless --force is given, and the archive/,
source/ and template/ folders are always refused as output locations.
Requires python-pptx (which installs lxml).
"""
import argparse
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

PROJECT = Path(__file__).resolve().parent.parent
DECK_NAME = "GIA_HK_2027_Contingent_Worker_Business_Case_v4"
PROTECTED = [PROJECT / d for d in ("archive", "source", "template")]

ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
ap.add_argument("--template", type=Path, default=PROJECT / "template" / "aia_ppt_template.pptx")
ap.add_argument("--output", type=Path, default=PROJECT / "current" / f"{DECK_NAME}.pptx")
ap.add_argument("--force", action="store_true", help="overwrite the output file if it already exists")
args = ap.parse_args()

TEMPLATE = args.template
OUT = args.output.resolve()
if any(OUT.is_relative_to(p.resolve()) for p in PROTECTED):
    sys.exit(f"Refusing to write into a protected folder: {OUT}")
if OUT.exists() and not args.force:
    sys.exit(f"{OUT} already exists. Use --force to overwrite, or --output for a new path.")
OUT.parent.mkdir(parents=True, exist_ok=True)

RED = RGBColor(0xD3, 0x11, 0x45)
CHAR = RGBColor(0x33, 0x3D, 0x47)
GREY_TXT = RGBColor(0x5F, 0x67, 0x70)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
TINT = RGBColor(0xFB, 0xE7, 0xEC)
LINE = "BFBFBF"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation(TEMPLATE)

# --- sensitivity marking: keep the template's label text unchanged (the label itself must be
# applied in PowerPoint), but stop the marking box wrapping so it renders on one line
for sh in prs.slide_master.shapes:
    if sh.has_text_frame and sh._element.find(".//" + qn("p:nvPr")) is not None and \
            "classification" in etree.tostring(sh._element.find(".//" + qn("p:nvPr"))).decode():
        sh.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")
        sh.width = Inches(1.4)

# --- remove the template's sample slides (parts are dropped on save)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

L_TITLE, L_CONTENT = prs.slide_layouts[0], prs.slide_layouts[1]
SLDNUM_SP = [sh for sh in L_CONTENT.placeholders if sh.placeholder_format.idx == 4][0]._element


def ph(slide, idx):
    return [p for p in slide.placeholders if p.placeholder_format.idx == idx][0]


def set_runs(tf, text, size=None, bold=None, color=None):
    p = tf.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = text
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if color is not None: r.font.color.rgb = color
    r.font.name = FONT
    return r


def bullet(p, char="•", indent=0.18):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", FONT)
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", char)


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0"); pPr.set("indent", "0")
    etree.SubElement(pPr, qn("a:buNone"))


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0, fill=None):
    """paras: list of dicts {runs:[(text, {bold,color,size})], size, color, bold, bullet, after, before, align}"""
    if fill is not None:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid(); box.fill.fore_color.rgb = fill
        box.line.fill.background()
        box.shadow.inherit = False
        st = box._element.find(qn("p:style"))
        if st is not None:
            box._element.remove(st)
    else:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    m = Inches(margin)
    tf.margin_left = tf.margin_right = m
    tf.margin_top = tf.margin_bottom = Inches(min(margin, 0.12)) if margin else 0
    tf.vertical_anchor = anchor
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pd.get("bullet"):
            bullet(p)
        else:
            no_bullet(p)
        p.alignment = pd.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(pd.get("after", 0))
        p.space_before = Pt(pd.get("before", 0))
        p.line_spacing = pd.get("ls", 1.0)
        runs = pd.get("runs") or [(pd["t"], {})]
        for t, o in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT
            f.size = Pt(o.get("size", pd.get("size", 13)))
            f.bold = o.get("bold", pd.get("bold", False))
            f.italic = o.get("italic", pd.get("italic", False))
            f.color.rgb = o.get("color", pd.get("color", CHAR))
    return box


def content_slide(title, message):
    s = prs.slides.add_slide(L_CONTENT)
    for idx in (15, 16):
        e = ph(s, idx)._element; e.getparent().remove(e)
    s.shapes._spTree.append(copy.deepcopy(SLDNUM_SP))
    t = ph(s, 10)
    t.left, t.top, t.width, t.height = Inches(0.39), Inches(0.55), Inches(12.2), Inches(0.62)
    t.text_frame.margin_left = 0
    set_runs(t.text_frame, title, size=28, bold=True, color=CHAR)
    text(s, 0.39, 1.22, 12.55, 0.5, [dict(t=message, size=16, color=RED)])
    return s


def footnote(s, t, y=6.22):
    text(s, 0.39, y, 11.5, 0.3, [dict(t=t, size=10, color=GREY_TXT, italic=True)])


# ---------- table helpers
def cell_border(cell, sides="B", color=LINE, w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        tag = qn(f"a:ln{side}")
        for e in tcPr.findall(tag):
            tcPr.remove(e)
    for side in ("L", "R", "T", "B"):
        ln = etree.SubElement(tcPr, qn(f"a:ln{side}"))
        if side in sides:
            ln.set("w", str(w)); ln.set("cmpd", "sng")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
        else:
            ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
    # schema order: borders must precede fill
    fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f); tcPr.append(f)


def fill_cell(cell, paras, fill=None, anchor=MSO_ANCHOR.TOP):
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.07); cell.margin_bottom = Inches(0.07)
    cell.vertical_anchor = anchor
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pd.get("bullet"):
            bullet(p, indent=0.15)
        p.space_after = Pt(pd.get("after", 0))
        runs = pd.get("runs") or [(pd["t"], {})]
        for t, o in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT
            f.size = Pt(o.get("size", pd.get("size", 12)))
            f.bold = o.get("bold", pd.get("bold", False))
            f.color.rgb = o.get("color", pd.get("color", CHAR))
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill


def table(s, x, y, w, col_w, row_h):
    shp = s.shapes.add_table(len(row_h), len(col_w), Inches(x), Inches(y), Inches(w), Inches(sum(row_h)))
    tbl = shp.table
    # "No Style, No Grid" so only our explicit formatting applies
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for i, rh in enumerate(row_h):
        tbl.rows[i].height = Inches(rh)
    return tbl


# =====================================================================
# Slide 1 — Title
s = prs.slides.add_slide(L_TITLE)
set_runs(ph(s, 10).text_frame, "Business Case for 2027 Contingent Worker")
sub = ph(s, 1)
sub.left, sub.top, sub.width, sub.height = Inches(0.39), Inches(4.5), Inches(6.3), Inches(0.9)
tf = sub.text_frame
set_runs(tf, "GIA HK Data Analytics")
p = tf.add_paragraph()
r = p.add_run(); r.text = "Request for 1,400 hours of technical support in 2027"; r.font.name = FONT

# =====================================================================
# Slide 2 — Executive summary
s = content_slide("Executive Summary and Approval Request",
                  "Additional technical capacity will support 2027 audit delivery and practical automation within GIA HK.")

text(s, 0.39, 2.0, 3.45, 3.4, [
    dict(t="REQUEST", size=11, bold=True, color=GREY_TXT, after=6),
    dict(t="1,400 hours", size=40, bold=True, color=RED, after=4),
    dict(t="of technical support in 2027", size=14, after=16),
    dict(t="One contingent worker", size=14, bold=True, after=4),
    dict(t="GIA HK Data Analytics", size=14),
], margin=0.25, fill=LIGHT)

rows = [
    ("Why it is needed",
     "The three-person Data Analytics team is fully occupied with ongoing work. Planned 2027 engagements, "
     "automation, Group initiatives and tool maintenance all need technical support."),
    ("What it will deliver",
     "Technical support for planned audit engagements, and practical AI and automation solutions "
     "developed directly with business auditors."),
    ("Why it benefits GIA HK",
     "Technical capacity within the audit team, responsive to audit priorities. The existing team stays "
     "focused on its ongoing responsibilities, and TSS and DNA support continues."),
]
tbl = table(s, 4.14, 2.0, 8.8, [2.4, 6.4], [1.13] * 3)
for i, (k, v) in enumerate(rows):
    fill_cell(tbl.cell(i, 0), [dict(t=k, size=14, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    fill_cell(tbl.cell(i, 1), [dict(t=v, size=14)], anchor=MSO_ANCHOR.MIDDLE)
    for c in (0, 1):
        cell_border(tbl.cell(i, c), sides="TB" if i == 0 else "B")

text(s, 0.39, 5.62, 12.55, 0.5, [dict(runs=[
    ("Decision requested:  ", dict(bold=True, color=WHITE)),
    ("Approval of one contingent worker for GIA HK Data Analytics in 2027, covering up to 1,400 hours.",
     dict(color=WHITE))], size=14)], anchor=MSO_ANCHOR.MIDDLE, margin=0.2, fill=CHAR)

# =====================================================================
# Slide 3 — Workload and capacity
s = content_slide("2027 Audit Workload and Existing Capacity",
                  "Planned 2027 work needs technical support that the current team has limited capacity to absorb.")

hdr = ["Area", "Planned 2027 activity", "Technical support required"]
body = [
    ("Engagements", "Individual Life Claims Audit", None),
    ("", "Premium Collection Audit", None),
    ("", "Post-Sales Call Analytics", None),
    ("", "Complaint Handling", None),
    ("Automation", "Audit automation and data pipelines",
     "Automate data extraction and build data pipelines towards continuous auditing"),
    ("Group initiatives", "Group AotF and Databricks Genie Space",
     "HK support for the Group AotF workstream and the Genie Space pilot"),
    ("Platforms", "Maintenance of existing analytics platforms",
     "Maintain and enhance Databricks, SandBox and the 2026 tools"),
]
tbl = table(s, 0.39, 2.0, 8.75, [1.75, 3.2, 3.8], [0.42] + [0.38] * 4 + [0.58, 0.58, 0.58])
for c, h in enumerate(hdr):
    fill_cell(tbl.cell(0, c), [dict(t=h, size=12, bold=True, color=WHITE)], fill=CHAR, anchor=MSO_ANCHOR.MIDDLE)
    cell_border(tbl.cell(0, c), sides="")
for i, (a, act, sup) in enumerate(body, start=1):
    fill_cell(tbl.cell(i, 0), [dict(t=a, size=12, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    fill_cell(tbl.cell(i, 1), [dict(t=act, size=12)], anchor=MSO_ANCHOR.MIDDLE)
    if sup:
        fill_cell(tbl.cell(i, 2), [dict(t=sup, size=12)], anchor=MSO_ANCHOR.MIDDLE)
for c in range(3):
    for i in range(1, 8):
        cell_border(tbl.cell(i, c), sides="B" if (i >= 4 or c == 1) else "")
m = tbl.cell(1, 0); m.merge(tbl.cell(4, 0))
m2 = tbl.cell(1, 2); m2.merge(tbl.cell(4, 2))
fill_cell(m2, [dict(t="Technical and analytics support to the engagement teams", size=12)], anchor=MSO_ANCHOR.MIDDLE)
for c in (0, 2):
    cell_border(tbl.cell(1, c), sides="B")

text(s, 9.44, 2.0, 3.5, 3.9, [
    dict(t="EXISTING CAPACITY", size=11, bold=True, color=GREY_TXT, after=6),
    dict(t="3 employees", size=32, bold=True, color=RED, after=4),
    dict(t="in GIA HK Data Analytics", size=13, after=14),
    dict(t="Fully occupied with ongoing audit and analytics work", size=13, bullet=True, after=8),
    dict(t="Limited capacity for new development alongside the current workload", size=13, bullet=True, after=8),
    dict(t="Hours to be allocated by audit priority; engagement support remains a core responsibility",
         size=13, bullet=True),
], margin=0.25, fill=LIGHT)

footnote(s, "Engagements are planned for 2027 and subject to confirmation in the final audit plan. "
            "Team capacity reflects GIA HK management's assessment.")

# =====================================================================
# Slide 4 — 2026 capability as the foundation for 2027 value
s = content_slide("From 2026 Capability to 2027 Business Value",
                  "Technical capabilities demonstrated in 2026 provide a foundation for stronger audit delivery, "
                  "practical automation, and greater flexibility in 2027.")

gap = 0.305
solutions = [
    ("Regulatory Automation",
     "Retrieves regulatory documents and identifies potentially relevant audit requirements. "
     "Covers approx. 259 documents.*",
     "Essentially usable"),
    ("MedVerify",
     "Checks medical providers against official regulator listings to flag potentially suspicious cases.",
     "Developed; awaiting operational approval"),
    ("Welcome Call UI",
     "Lets business auditors run post-sales call analytics through a user-friendly interface.",
     "Front end essentially complete; backend integration outstanding"),
    ("Supporting tools",
     "SecureSheet (data masking), DataCheck (data quality checks), Relationship Mapper (dataset linkage)",
     "Expected to be usable; adoption not yet confirmed"),
]
TOP, HDR_H, ROW_H = 2.05, 0.42, 0.86
tbl = table(s, 0.39, TOP, 7.85, [1.95, 3.75, 2.15], [HDR_H] + [ROW_H] * len(solutions))
for c, h in enumerate(["2026 solution", "What it does", "Current status"]):
    fill_cell(tbl.cell(0, c), [dict(t=h, size=13, bold=True, color=WHITE)], fill=CHAR, anchor=MSO_ANCHOR.MIDDLE)
    cell_border(tbl.cell(0, c), sides="")
for i, (name, what, status) in enumerate(solutions, start=1):
    fill_cell(tbl.cell(i, 0), [dict(t=name, size=13, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    fill_cell(tbl.cell(i, 1), [dict(t=what, size=12)], anchor=MSO_ANCHOR.MIDDLE)
    fill_cell(tbl.cell(i, 2), [dict(t=status, size=12, bold=True, color=RED)], anchor=MSO_ANCHOR.MIDDLE)
    for c in range(3):
        cell_border(tbl.cell(i, c), sides="B")

bottom = TOP + HDR_H + ROW_H * len(solutions)
arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(8.4), Inches((TOP + bottom) / 2 - 0.25),
                           Inches(0.3), Inches(0.5))
arrow.fill.solid(); arrow.fill.fore_color.rgb = RED
arrow.line.fill.background()
st = arrow._element.find(qn("p:style"))
if st is not None:
    arrow._element.remove(st)

RX, RW = 8.86, 12.94 - 8.86
text(s, RX, TOP, RW, HDR_H, [dict(t="Expected business value in 2027", size=13, bold=True, color=WHITE)],
     anchor=MSO_ANCHOR.MIDDLE, margin=0.1, fill=RED)
benefits = [
    ("Audit delivery", "Technical and analytics support for planned 2027 audit engagements."),
    ("Practical AI and automation",
     "Solutions developed directly with business auditors to help reduce repetitive manual work."),
    ("Technical flexibility", "Responds to changing audit priorities while existing tools remain maintained."),
]
bgap = 0.1
bh = (bottom - (TOP + HDR_H) - bgap * len(benefits)) / len(benefits)
for i, (head, desc) in enumerate(benefits):
    text(s, RX, TOP + HDR_H + bgap + i * (bh + bgap), RW, bh, [
        dict(t=head, size=14, bold=True, after=4),
        dict(t=desc, size=12),
    ], anchor=MSO_ANCHOR.MIDDLE, margin=0.18, fill=TINT)

footnote(s, "Tool adoption and operational status are subject to confirmation; time savings have not yet been "
            "measured. *Document count is subject to verification. 2027 benefits are expected, not measured "
            "results; proposed measures are set out on slide 6.", y=6.05)
s.shapes[-1].width = Inches(12.55)

# =====================================================================
# Slide 5 — Resource model
s = content_slide("Why Dedicated Technical Capacity Within GIA HK?",
                  "GIA HK benefits from technical expertise that works directly with its auditors on its priorities.")

hdr = ["", "Existing DA team", "TSS and DNA", "Contingent worker (proposed)"]
rows = [
    ("Main focus", "Ongoing audit and analytics work",
     "Technical development for GIA HK and other areas", "Engagement support and development within GIA HK"),
    ("Availability for GIA HK development", "Limited; team fully occupied",
     "Scheduled alongside their other priorities", "Dedicated to GIA HK audit priorities"),
    ("Working with business auditors", "Direct, within existing workload",
     "Through cross-team planning", "Direct and day-to-day"),
]
tbl = table(s, 0.39, 2.0, 12.55, [2.65, 3.3, 3.3, 3.3], [0.48, 0.72, 0.72, 0.72])
for c, h in enumerate(hdr):
    fill_cell(tbl.cell(0, c), [dict(t=h, size=13, bold=True, color=WHITE)],
              fill=CHAR if c < 3 else RED, anchor=MSO_ANCHOR.MIDDLE)
    cell_border(tbl.cell(0, c), sides="")
for i, row in enumerate(rows, start=1):
    for c, v in enumerate(row):
        fill_cell(tbl.cell(i, c), [dict(t=v, size=13, bold=(c == 0))],
                  fill=TINT if c == 3 else None, anchor=MSO_ANCHOR.MIDDLE)
        cell_border(tbl.cell(i, c), sides="B")

text(s, 0.39, 4.95, 12.55, 1.1, [
    dict(t="Why a contingent worker", size=14, bold=True, after=6),
    dict(t="Adds dedicated technical capacity for a defined period (2027, up to 1,400 hours). Existing "
           "familiarity with GIA processes and technology can reduce onboarding. TSS and DNA continue to "
           "support GIA HK where appropriate.", size=13),
], margin=0.22, fill=LIGHT)

# =====================================================================
# Slide 6 — Recommendation
s = content_slide("Recommendation, Governance and Approval Request",
                  "A defined one-year resource, allocated by audit priority and reviewed against practical measures.")

half = (12.55 - gap) / 2
text(s, 0.39, 2.0, half, 2.75, [
    dict(t="How the resource will be managed", size=15, bold=True, after=10),
    dict(t="Allocated by audit priority, under the direction of the Data Analytics lead",
         size=13, bullet=True, after=8),
    dict(t="Supports both audit engagements and technical development; engagement support remains core",
         size=13, bullet=True, after=8),
    dict(t="Development work coordinated with TSS and DNA where relevant", size=13, bullet=True),
], margin=0.22, fill=LIGHT)

text(s, 0.39 + half + gap, 2.0, half, 2.75, [
    dict(t="Proposed accountability measures", size=15, bold=True, after=10),
    dict(t="Audit engagements supported", size=13, bullet=True, after=6),
    dict(t="Automation solutions developed or improved", size=13, bullet=True, after=6),
    dict(t="Solutions adopted by auditors", size=13, bullet=True, after=6),
    dict(t="Reduction in manual effort, where measurable", size=13, bullet=True, after=12),
    dict(t="A management progress review during 2027 is proposed.", size=13, italic=True),
], margin=0.22, fill=LIGHT)

text(s, 0.39, 5.05, 12.55, 0.55, [dict(runs=[
    ("Decision requested:  ", dict(bold=True, color=WHITE)),
    ("Approve one contingent worker for GIA HK Data Analytics for 2027, covering up to 1,400 hours.",
     dict(color=WHITE))], size=14)], anchor=MSO_ANCHOR.MIDDLE, margin=0.2, fill=CHAR)

footnote(s, "Cost and contract terms will be assessed separately by HR and Finance.", y=5.75)

prs.save(OUT)
print(f"Saved {OUT}")
