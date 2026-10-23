"""Check the generated Version 3 deck against the benchmark results, and dump its text for review.

    python3 ppt/ocr_benchmark/scripts/check_deck_v3.py      (after generate_ppt_v3.py)

Checks:
- 12 slides, every slide has speaker notes;
- PaddleOCR Heavy and its accuracy figures appear nowhere (slides or notes);
- the initial results (original rules 37.3% / 41.2%, AOAI 88.2%) appear, and a slide showing an original
  rule figure also explains the missing Chinese support;
- the development results (improved rules 88.2% / 86.3%, AOAI prompt v2 96.1%) appear, and a slide showing
  86.3% or 96.1% also calls them development results;
- the AOAI time appears.
Writes previews/deck_text.md (all slide text and notes). Exits non-zero on failure.
"""
import csv
import sys
from pathlib import Path

from pptx import Presentation

PROJECT = Path(__file__).resolve().parent.parent
RES = PROJECT.parent.parent / "python" / "receipt_ocr_benchmark" / "results" / "v2"
DECK = PROJECT / "current" / "OCR_Benchmark_V3.pptx"


def summary(path):
    with open(path, newline="") as fh:
        return {r["approach"]: r for r in csv.DictReader(fh)}


SUM, SUM_V2 = summary(RES / "scores_improved" / "summary.csv"), summary(RES / "scores_aoai_improved" / "summary.csv")
pct = lambda r: f"{float(r['accuracy_pct']):.1f}%"  # noqa: E731
AOAI_PCT, AOAI_V2_PCT = pct(SUM["azure_di_aoai"]), pct(SUM_V2["aoai_v2_rows"])
AOAI_SECS = f"{int(SUM['azure_di_aoai']['mean_ms_per_receipt']) / 1000:.1f} s"
INITIAL = {pct(SUM[a]) for a in ("paddle_light_general", "azure_di_general")}
DEVELOPMENT = {pct(SUM[a]) for a in ("paddle_light_improved", "azure_di_improved")} | {AOAI_V2_PCT}
HEAVY = {pct(r) for a, r in SUM.items() if "heavy" in a} - INITIAL - DEVELOPMENT - {AOAI_PCT}

prs = Presentation(DECK)
problems, out = [], [f"# {DECK.name}: text and notes\n"]
if len(prs.slides) != 12:
    problems.append(f"expected 12 slides, found {len(prs.slides)}")
all_slide_text = []
for i, s in enumerate(prs.slides, 1):
    out.append(f"\n## Slide {i}\n")
    paras = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append("- " + sh.text_frame.text.replace("\n", " / "))
            paras += [p.text for p in sh.text_frame.paragraphs]
        if sh.has_table:
            for row in sh.table.rows:
                cells = [c.text.replace("\n", " ") for c in row.cells]
                out.append("  | " + " | ".join(cells) + " |")
                paras += cells
    notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    out.append(f"\n  Notes: {notes}")
    slide_text = "\n".join(paras)
    all_slide_text.append(slide_text)
    if not notes.strip():
        problems.append(f"slide {i}: no speaker notes")
    for where, t in (("slide", slide_text), ("notes", notes)):
        if "heavy" in t.lower() or any(p in t for p in HEAVY):
            problems.append(f"slide {i} {where}: mentions PaddleOCR Heavy or its results")
    low = slide_text.lower()
    if any(p in slide_text for p in INITIAL) and "chinese" not in low:
        problems.append(f"slide {i}: original rule accuracy without the missing Chinese support explained")
    if any(p in slide_text for p in DEVELOPMENT - {AOAI_PCT}) and "development" not in low:
        problems.append(f"slide {i}: development result without a development label")

joined = "\n".join(all_slide_text)
# Accuracy by field: every "pct  (correct/applicable)" of the three tuned approaches, counted from detail.csv,
# on one slide that labels them as development results.
with open(RES / "scores_aoai_improved" / "detail.csv", newline="") as fh:
    detail = [r for r in csv.DictReader(fh) if r["result"] in ("Correct", "Wrong")]
by_field = [s for s in all_slide_text if "Accuracy by Field" in s]
if len(by_field) != 1:
    problems.append("expected one 'Accuracy by Field' slide")
else:
    if "development" not in by_field[0].lower():
        problems.append("Accuracy by Field slide is not labelled as development results")
    for a in ("paddle_light_improved", "azure_di_improved", "aoai_v2_rows"):
        for f in {r["field"] for r in detail}:
            rows = [r for r in detail if r["approach"] == a and r["field"] == f]
            c, n = sum(r["result"] == "Correct" for r in rows), len(rows)
            if f"{100 * c / n:.1f}%  ({c}/{n})" not in by_field[0]:
                problems.append(f"Accuracy by Field: {a} {f} {c}/{n} not shown")
for needed in sorted(INITIAL | DEVELOPMENT | {AOAI_PCT, AOAI_SECS}):
    if needed not in joined:
        problems.append(f"{needed} not found in slide text")

(PROJECT / "previews").mkdir(exist_ok=True)
(PROJECT / "previews" / "deck_text.md").write_text("\n".join(out) + "\n")
if problems:
    sys.exit("Deck check FAILED:\n- " + "\n- ".join(problems))
print(f"Deck check passed: {len(prs.slides)} slides; no Heavy; initial {sorted(INITIAL | {AOAI_PCT})} and "
      f"development {sorted(DEVELOPMENT)} results present and labelled; text in previews/deck_text.md")
