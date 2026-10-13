"""Generate the Version 2 Receipt OCR Benchmark deck (8 slides) from the AIA template.

    python3 ppt/ocr_benchmark/scripts/generate_ppt_v2.py --force     (from the repository root)

Reads committed results only (no OCR, scoring or Azure calls):
- python/receipt_ocr_benchmark/results/v2/scores_improved/summary.csv and detail.csv: every
  accuracy and time figure, and every value shown in the receipt examples;
- results/v2/<approach dir>/receiptN.json: each example value is cross-checked against these files,
  and the build stops if they disagree with detail.csv;
- results/v2/paddle_light and azure_di raw OCR: statements that "the OCR read the value" are
  checked against them;
- results/v2/azure_di/receiptN.json text positions: used to crop and highlight the receipt photos
  (data/receipts_v2/benchmark_inputs, turned upright via EXIF as the benchmark inputs were);
- receipt_bench/extraction_general.py and extraction_improved.py: label vocabulary sizes (constants
  only; no extraction is run).

Receipt crops go to ppt/ocr_benchmark/assets_v2/. Never writes into archive/ or template/.
"""
import argparse
import copy
import csv
import json
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT = Path(__file__).resolve().parent.parent
REPO = PROJECT.parent.parent
BENCH = REPO / "python" / "receipt_ocr_benchmark"
RES = BENCH / "results" / "v2"
ASSETS = PROJECT / "assets_v2"
PROTECTED = [PROJECT / d for d in ("archive", "template")]

ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
ap.add_argument("--template", type=Path, default=PROJECT / "template" / "aia_ppt_template.pptx")
ap.add_argument("--output", type=Path, default=PROJECT / "current" / "OCR_Benchmark_V2.pptx")
ap.add_argument("--force", action="store_true", help="overwrite the output file if it already exists")
args = ap.parse_args()
OUT = args.output.resolve()
if any(OUT.is_relative_to(p.resolve()) for p in PROTECTED):
    sys.exit(f"Refusing to write into a protected folder: {OUT}")
if OUT.exists() and not args.force:
    sys.exit(f"{OUT} already exists. Use --force to overwrite, or --output for a new path.")
OUT.parent.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(exist_ok=True)

# ===================================================================== benchmark data (source of truth)
ORIG = {"light": "paddle_light_general", "heavy": "paddle_heavy_general", "di": "azure_di_general"}
IMPR = {"light": "paddle_light_improved", "heavy": "paddle_heavy_improved", "di": "azure_di_improved"}
AOAI = "azure_di_aoai"
RESULT_DIR = {"paddle_light_general": "extracted_general_light", "paddle_light_improved": "extracted_improved_light",
              "paddle_heavy_general": "extracted_general_heavy", "paddle_heavy_improved": "extracted_improved_heavy",
              "azure_di_general": "extracted_general_azure_di", "azure_di_improved": "extracted_improved_azure_di",
              "azure_di_aoai": "di_aoai"}
FIELD_NAMES = {"doctor_name": "Doctor", "registration_number": "Reg. no.", "patient_name": "Patient",
               "receipt_number": "Receipt no.", "service_date": "Service date", "total_amount": "Total"}
FIELDS = list(FIELD_NAMES)

with open(RES / "scores_improved" / "summary.csv", newline="") as fh:
    SUM = {r["approach"]: r for r in csv.DictReader(fh)}
with open(RES / "scores_improved" / "detail.csv", newline="") as fh:
    DETAIL = {(r["approach"], r["receipt"], r["field"]): r for r in csv.DictReader(fh)}
assert set(SUM) == set(RESULT_DIR), f"unexpected approaches in summary.csv: {sorted(SUM)}"
RECEIPTS = sorted({k[1] for k in DETAIL}, key=lambda s: int(s.replace("receipt", "")))
N_RECEIPTS, APPLICABLE = len(RECEIPTS), int(SUM[AOAI]["applicable"])
assert all(int(SUM[a]["applicable"]) == APPLICABLE for a in SUM)


def acc(a):
    return f"{float(SUM[a]['accuracy_pct']):.1f}%"


def count(a):
    return f"{int(SUM[a]['correct'])}/{APPLICABLE}"


def secs(a):
    return f"{int(SUM[a]['mean_ms_per_receipt']) / 1000:.1f} s"


def receipt_score(a, receipt):
    rows = [DETAIL[(a, receipt, f)] for f in FIELDS if DETAIL[(a, receipt, f)]["result"] in ("Correct", "Wrong")]
    return sum(r["result"] == "Correct" for r in rows), len(rows)


def json_value(approach, receipt, field):
    v = json.loads((RES / RESULT_DIR[approach] / f"{receipt}.json").read_text()).get(field)
    return "" if v is None else str(v)


def example(approach, receipt, field):
    """(value as extracted, correct?) from detail.csv, cross-checked against the extraction JSON."""
    row = DETAIL[(approach, receipt, field)]
    if json_value(approach, receipt, field) != row["predicted"]:
        sys.exit(f"Mismatch {approach} {receipt} {field}: detail.csv vs extraction JSON")
    return (row["predicted"] or "not found"), row["result"] == "Correct"


def truth(receipt, field):
    return DETAIL[(AOAI, receipt, field)]["ground_truth"]


def ocr_read(receipt, value, engines=("paddle_light", "azure_di")):
    """Stop the build unless every engine's raw OCR text contains `value` (backs 'the OCR read it')."""
    for e in engines:
        if value not in (RES / e / f"{receipt}.json").read_text():
            sys.exit(f"{e} OCR of {receipt} does not contain {value!r}; the slide statement would be wrong")


TOKENS_IN = round(int(SUM[AOAI]["aoai_input_tokens"]) / N_RECEIPTS)
TOKENS_OUT = round(int(SUM[AOAI]["aoai_output_tokens"]) / N_RECEIPTS)

# Label vocabulary sizes, counted from the two rule modules (constants only).
sys.path.insert(0, str(BENCH))
from receipt_bench import extraction_general as G  # noqa: E402
from receipt_bench import extraction_improved as I  # noqa: E402

VOCAB_ORIG = len({l for group in (G.DOCTOR_LABELS, G.PATIENT_LABELS, G.REG_LABELS, G.RECEIPT_LABELS,
                                  G.TOTAL_LABELS, G.PAID_LABELS, G.SUBTOTAL_LABELS, G.CHARGE_LABELS,
                                  G.BALANCE_LABELS) for l in group} | {l for g in G.DATE_LABELS for l in g})
VOCAB_IMPR = len({l for en, zh in I.LABELS.values() for l in en + zh})

# ===================================================================== receipt crops
INPUTS = BENCH / "data" / "receipts_v2" / "benchmark_inputs"


def receipt_image(receipt):
    src = next(p for p in INPUTS.iterdir() if p.stem == receipt).resolve()
    return ImageOps.exif_transpose(Image.open(src)).convert("RGB")


def find_boxes(receipt, needles):
    page = json.loads((RES / "azure_di" / f"{receipt}.json").read_text())["pages"][0]
    blocks = [(b["text"], min(p[0] for p in b["polygon"]), min(p[1] for p in b["polygon"]),
               max(p[0] for p in b["polygon"]), max(p[1] for p in b["polygon"])) for b in page["blocks"]]
    boxes = []
    for n in needles:
        hit = next((b for b in blocks if n in b[0]), None)
        if hit is None:
            sys.exit(f"{receipt}: text {n!r} not found in Azure DI output; cannot place the crop")
        boxes.append(hit[1:])
    return (page["width"], page["height"]), boxes


def make_crop(receipt, needles, name, highlight, max_px=2000):
    img = receipt_image(receipt)
    (w, h), boxes = find_boxes(receipt, needles)
    if img.size != (w, h):
        sys.exit(f"{receipt}: image {img.size} does not match the Azure DI page {(w, h)}")
    m = int(0.025 * w)
    x0, y0 = max(0, min(b[0] for b in boxes) - m), max(0, min(b[1] for b in boxes) - m)
    x1, y1 = min(w, max(b[2] for b in boxes) + m), min(h, max(b[3] for b in boxes) + m)
    draw = ImageDraw.Draw(img)
    lw = max(3, w // 500)
    for b in find_boxes(receipt, highlight)[1]:
        draw.rectangle([b[0] - lw * 2, b[1] - lw * 2, b[2] + lw * 2, b[3] + lw * 2], outline=(211, 17, 69), width=lw)
    crop = img.crop((x0, y0, x1, y1))
    crop.thumbnail((max_px, max_px))
    path = ASSETS / f"{name}.jpg"
    crop.save(path, quality=90)
    return path


def make_thumb(receipt, name, max_px=1000):
    img = receipt_image(receipt)
    img.thumbnail((max_px, max_px))
    path = ASSETS / f"{name}.jpg"
    img.save(path, quality=88)
    return path


# ===================================================================== style (AIA template, plain)
RED = RGBColor(0xD3, 0x11, 0x45)
CHAR = RGBColor(0x33, 0x3D, 0x47)
GREY_TXT = RGBColor(0x5F, 0x67, 0x70)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
GOOD = RGBColor(0x1E, 0x7B, 0x4A)
GOOD_TINT = RGBColor(0xE6, 0xF2, 0xEA)
BAD_TINT = RGBColor(0xFB, 0xE7, 0xEC)
LINE = "BFBFBF"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT, FONT_ZH = "Arial", "Microsoft JhengHei"

prs = Presentation(args.template)
for sh in prs.slide_master.shapes:  # keep the classification marking on one line (label applied in PowerPoint)
    nv = sh._element.find(".//" + qn("p:nvPr"))
    if sh.has_text_frame and nv is not None and "classification" in etree.tostring(nv).decode():
        sh.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")
        sh.width = Inches(1.4)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
L_TITLE, L_CONTENT = prs.slide_layouts[0], prs.slide_layouts[1]
SLDNUM_SP = [sh for sh in L_CONTENT.placeholders if sh.placeholder_format.idx == 4][0]._element


def ph(slide, idx):
    return [p for p in slide.placeholders if p.placeholder_format.idx == idx][0]


def style_run(r, size=None, bold=None, color=None, italic=None):
    f = r.font
    f.name = FONT
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn("a:ea")):
        rPr.remove(e)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", FONT_ZH)


def set_runs(tf, text_, size=None, bold=None, color=None):
    p = tf.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = text_
    style_run(r, size, bold, color)


def bullet(p, indent=0.2):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0"); pPr.set("indent", "0")
    etree.SubElement(pPr, qn("a:buNone"))


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0, fill=None):
    if fill is not None:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid(); box.fill.fore_color.rgb = fill
        box.line.fill.background()
        st = box._element.find(qn("p:style"))
        if st is not None:
            box._element.remove(st)
    else:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(min(margin, 0.12)) if margin else 0
    tf.vertical_anchor = anchor
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bullet(p) if pd.get("bullet") else no_bullet(p)
        p.alignment = pd.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(pd.get("after", 0))
        for t, o in pd.get("runs") or [(pd["t"], {})]:
            r = p.add_run(); r.text = t
            style_run(r, o.get("size", pd.get("size", 16)), o.get("bold", pd.get("bold", False)),
                      o.get("color", pd.get("color", CHAR)), o.get("italic", pd.get("italic", False)))
    return box


def content_slide(title, message, notes):
    s = prs.slides.add_slide(L_CONTENT)
    for idx in (15, 16):
        e = ph(s, idx)._element; e.getparent().remove(e)
    s.shapes._spTree.append(copy.deepcopy(SLDNUM_SP))
    t = ph(s, 10)
    t.left, t.top, t.width, t.height = Inches(0.39), Inches(0.55), Inches(12.2), Inches(0.62)
    t.text_frame.margin_left = 0
    set_runs(t.text_frame, title, size=28, bold=True, color=CHAR)
    text(s, 0.39, 1.22, 12.55, 0.5, [dict(t=message, size=17, color=RED)])
    s.notes_slide.notes_text_frame.text = notes
    return s


def note(s, t, y=6.15):
    text(s, 0.39, y, 12.0, 0.35, [dict(t=t, size=11, color=GREY_TXT, italic=True)])


def cell_border(cell, sides="B", color=LINE, w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        for e in tcPr.findall(qn(f"a:ln{side}")):
            tcPr.remove(e)
    for side in ("L", "R", "T", "B"):
        ln = etree.SubElement(tcPr, qn(f"a:ln{side}"))
        if side in sides:
            ln.set("w", str(w)); ln.set("cmpd", "sng")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        else:
            ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
    for f in [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]:
        tcPr.remove(f); tcPr.append(f)


def fill_cell(cell, paras, fill=None, anchor=MSO_ANCHOR.MIDDLE):
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Inches(0.1)
    cell.margin_top = cell.margin_bottom = Inches(0.05)
    cell.vertical_anchor = anchor
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pd.get("align", PP_ALIGN.LEFT)
        for t, o in pd.get("runs") or [(pd["t"], {})]:
            r = p.add_run(); r.text = t
            style_run(r, o.get("size", pd.get("size", 14)), o.get("bold", pd.get("bold", False)),
                      o.get("color", pd.get("color", CHAR)), o.get("italic", pd.get("italic", False)))
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill


def table(s, x, y, col_w, row_h):
    tbl = s.shapes.add_table(len(row_h), len(col_w), Inches(x), Inches(y), Inches(sum(col_w)),
                             Inches(sum(row_h))).table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for i, rh in enumerate(row_h):
        tbl.rows[i].height = Inches(rh)
    return tbl


def header_row(tbl, heads, size=13, align=None):
    for c, hd in enumerate(heads):
        fill_cell(tbl.cell(0, c), [dict(t=hd, size=size, bold=True, color=WHITE,
                                        align=(align[c] if align else PP_ALIGN.LEFT))], fill=CHAR)
        cell_border(tbl.cell(0, c), sides="")


def picture(s, path, x, y, max_w, max_h):
    with Image.open(path) as im:
        ar = im.width / im.height
    w, h = (max_w, max_w / ar) if max_w / ar <= max_h else (max_h * ar, max_h)
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    pic.line.width = Pt(0.75)
    return x, y, w, h


def example_table(s, x, y, receipt, fields, col_w, row_h=0.72, size=14):
    """Correct value vs original rules vs improved rules (PaddleOCR Light) vs DI + AOAI."""
    tbl = table(s, x, y, list(col_w), [0.62] + [row_h] * len(fields))
    header_row(tbl, ["Field", "Correct", "Original rules", "Improved rules", "DI + AOAI"], size=12.5)
    for i, f in enumerate(fields, start=1):
        fill_cell(tbl.cell(i, 0), [dict(t=FIELD_NAMES[f], size=size, bold=True)])
        fill_cell(tbl.cell(i, 1), [dict(t=truth(receipt, f), size=size)])
        for c, a in enumerate((ORIG["light"], IMPR["light"], AOAI), start=2):
            v, ok = example(a, receipt, f)
            fill_cell(tbl.cell(i, c), [dict(runs=[("✓ " if ok else "✗ ", dict(bold=True, color=GOOD if ok else RED)),
                                                  (v, dict(color=CHAR if v != "not found" else GREY_TXT))],
                                            size=size)], fill=GOOD_TINT if ok else BAD_TINT)
        for c in range(len(col_w)):
            cell_border(tbl.cell(i, c), sides="B", color="FFFFFF", w=19050)
    return tbl


SOURCE = "Source: python/receipt_ocr_benchmark/results/v2/scores_improved (summary.csv, detail.csv) and extraction JSON."
TUNED = f"Improved rules were tuned using these {N_RECEIPTS} receipts. The AOAI prompt was not changed."

# ===================================================================== 1. title
s = prs.slides.add_slide(L_TITLE)
set_runs(ph(s, 10).text_frame, "Medical Receipt OCR Benchmark")
sub = ph(s, 1)
sub.left, sub.top, sub.width, sub.height = Inches(0.39), Inches(4.5), Inches(6.6), Inches(0.9)
set_runs(sub.text_frame, "Can open-source OCR with rules match Azure AI at extracting receipt fields?")
p = sub.text_frame.add_paragraph()
r = p.add_run(); r.text = f"{N_RECEIPTS} receipts · 6 fields · 4 approaches"; style_run(r)
s.notes_slide.notes_text_frame.text = (
    "Objective: compare open-source OCR plus extraction rules with Azure Document Intelligence plus Azure OpenAI "
    "for extracting six fields from medical receipts, on accuracy, speed, cost drivers and maintenance. " + SOURCE)

# ===================================================================== 2. what we tested
s = content_slide(
    "What We Tested",
    f"{N_RECEIPTS} receipts with different layouts and languages. Each approach reads the text, then finds 6 fields.",
    f"Receipts: English, bilingual and Traditional Chinese layouts. Fields: doctor, registration number, patient, "
    f"receipt number, service date, total. {APPLICABLE} fields were present on the receipts and scored. "
    "Step 1 (text recognition, OCR) turns the photo into text; step 2 (field extraction) picks the six values "
    "from that text. " + SOURCE)
for i, (rcpt, label) in enumerate((("receipt1", "English"), ("receipt3", "Bilingual"), ("receipt4", "Chinese"))):
    x, y, w, h = picture(s, make_thumb(rcpt, f"thumb_{rcpt}"), 0.39 + i * 2.25, 1.95, 2.1, 3.7)
    text(s, x, y + h + 0.08, 2.1, 0.3, [dict(t=label, size=12, color=GREY_TXT)])
tbl = table(s, 7.25, 1.95, [2.15, 1.85, 1.69], [0.62, 0.68, 0.68, 0.68, 0.68])
header_row(tbl, ["Approach", "1. Read text", "2. Find fields"])
for i, row in enumerate([("PaddleOCR Light", "Open-source OCR", "Rules"),
                         ("PaddleOCR Heavy", "Open-source OCR", "Rules"),
                         ("Azure DI", "Azure OCR", "Rules"),
                         ("Azure DI + AOAI", "Azure OCR", "Azure OpenAI")], start=1):
    for c, v in enumerate(row):
        fill_cell(tbl.cell(i, c), [dict(t=v, size=14, bold=(c == 0))])
        cell_border(tbl.cell(i, c), sides="B")
text(s, 7.25, 5.45, 5.69, 0.8, [dict(t=f"Fields: doctor, registration no., patient, receipt no., service date, total. "
                                       f"{APPLICABLE} fields scored.", size=13, color=GREY_TXT)])

# ===================================================================== 3. main results
s = content_slide(
    "Results: Field Accuracy and Speed",
    f"Improving the rules raised PaddleOCR Light from {acc(ORIG['light'])} to {acc(IMPR['light'])}. "
    f"AOAI reached {acc(AOAI)} without changes.",
    f"Accuracy = correct fields / {APPLICABLE} fields present on the {N_RECEIPTS} receipts. "
    f"Original and improved rules read the same OCR text. {TUNED} Time = average per receipt (PaddleOCR: CPU on "
    "Azure ML compute; Azure: API call time incl. network). " + SOURCE)
tbl = table(s, 0.39, 1.95, [3.6, 2.85, 3.25, 2.85], [0.72, 0.82, 0.82, 0.82, 0.82])
C = PP_ALIGN.CENTER
header_row(tbl, ["Approach", "Original rules", "Improved rules*", "Time per receipt"], size=16,
           align=[PP_ALIGN.LEFT, C, C, C])
rows = [("PaddleOCR Light", ORIG["light"], IMPR["light"], True),
        ("PaddleOCR Heavy", ORIG["heavy"], IMPR["heavy"], False),
        ("Azure DI", ORIG["di"], IMPR["di"], True)]
for i, (name, o, im, strong) in enumerate(rows, start=1):
    fill_cell(tbl.cell(i, 0), [dict(t=name, size=18, bold=True)])
    fill_cell(tbl.cell(i, 1), [dict(runs=[(acc(o), dict(size=20)), (f"  ({count(o)})", dict(size=13, color=GREY_TXT))],
                                    align=C)])
    fill_cell(tbl.cell(i, 2), [dict(runs=[(acc(im), dict(size=20, bold=strong, color=RED if strong else CHAR)),
                                          (f"  ({count(im)})", dict(size=13, color=GREY_TXT))], align=C)])
    fill_cell(tbl.cell(i, 3), [dict(t=secs(im), size=20, align=C)])
fill_cell(tbl.cell(4, 0), [dict(t="Azure DI + AOAI", size=18, bold=True)], fill=LIGHT)
fill_cell(tbl.cell(4, 1), [dict(runs=[(acc(AOAI), dict(size=20, bold=True, color=RED)),
                                      (f"  ({count(AOAI)})", dict(size=13, color=GREY_TXT))], align=C)], fill=LIGHT)
fill_cell(tbl.cell(4, 2), [dict(t="Not applicable (prompt unchanged)", size=13, color=GREY_TXT, align=C)], fill=LIGHT)
fill_cell(tbl.cell(4, 3), [dict(t=secs(AOAI), size=20, align=C)], fill=LIGHT)
for i in range(1, 5):
    for c in range(4):
        cell_border(tbl.cell(i, c), sides="B")
note(s, f"* Tuned using these {N_RECEIPTS} receipts, so not a measure of accuracy on new receipts. "
        f"Accuracy = correct fields out of {APPLICABLE}. All approaches processed all {N_RECEIPTS} receipts.", y=5.95)

# ===================================================================== 4. why improving the rules helped
for v in ("蔡嘉穎", "Dr. Siu Ka Wai"):
    ocr_read("receipt8", v)
s = content_slide(
    "Why Accuracy Improved: The Rules, Not the OCR",
    "The OCR read the text correctly. The original rules picked the wrong text.",
    "Receipt 8 (medical centre). Both OCR engines read the patient name and the doctor's name (checked in the raw "
    "OCR). The original rules matched 'Patient' and took the Chinese half of the label, 病人, as the name. The "
    "improved rules recognise bilingual labels. AOAI read the fields correctly with its unchanged prompt. Values: "
    "rules on PaddleOCR Light text; detail.csv, checked against the extraction JSON. " + SOURCE)
crop = make_crop("receipt8", ["Patient 病人", "蔡嘉穎", "Attending Doctor", "Dr. Siu Ka Wai"], "crop_receipt8",
                 ["Patient 病人", "蔡嘉穎", "Attending Doctor", "Dr. Siu Ka Wai"])
picture(s, crop, 0.39, 1.95, 5.6, 3.9)
example_table(s, 6.3, 1.95, "receipt8", ["patient_name", "doctor_name"], col_w=(1.15, 1.4, 1.35, 1.35, 1.39))
text(s, 6.3, 4.25, 6.64, 1.6, [
    dict(t="Original rules: took the Chinese label 病人 (\"patient\") as the patient's name.", size=14, bullet=True,
         after=8),
    dict(t="Improved rules: added bilingual labels, so the label and the value are kept apart.", size=14,
         bullet=True),
])
note(s, "Bilingual medical centre receipt. Rules shown on PaddleOCR Light text; Azure DI with the original rules "
        "made the same patient error.")

# ===================================================================== 5. the maintenance problem
for v in ("麥詠琳", "葉卓謙", "2026年6月9日"):
    ocr_read("receipt4", v)
o4, i4 = receipt_score(ORIG["light"], "receipt4"), receipt_score(IMPR["light"], "receipt4")
reg_o, ok_o = example(ORIG["light"], "receipt9", "receipt_number")
reg_i, ok_i = example(IMPR["light"], "receipt9", "receipt_number")
assert ok_o and not ok_i, "expected the receipt9 receipt-number regression"
s = content_slide(
    "The Trade-off: Rules Need Manual Maintenance",
    "High accuracy came from rewriting the rules by hand, and one fix broke another receipt.",
    f"Receipt 4 (Chinese medicine): the OCR read every value, but the original rules only knew English labels "
    f"({o4[0]} of {o4[1]} fields). After adding Chinese labels: {i4[0]} of {i4[1]}. The label vocabulary grew from "
    f"{VOCAB_ORIG} to {VOCAB_IMPR} phrases (counted in extraction_general.py and extraction_improved.py). On the vet "
    f"invoice (receipt 9) the original rules found receipt number {reg_o}; after the changes they return {reg_i}. "
    "Whether AOAI needs less maintenance on new layouts has not been tested yet. " + SOURCE)
crop = make_crop("receipt4", ["病人姓名", "麥詠琳", "診症日期", "2026年6月9日", "主診中醫師", "葉卓謙"],
                 "crop_receipt4", ["麥詠琳", "2026年6月9日", "葉卓謙"])
picture(s, crop, 0.39, 1.95, 4.6, 2.75)
text(s, 0.39, 4.8, 4.6, 0.5, [dict(t=f"Chinese medicine receipt: original rules {o4[0]}/{o4[1]} fields, "
                                     f"improved rules {i4[0]}/{i4[1]}.", size=12, color=GREY_TXT)])
for k, (big, small) in enumerate([(f"{VOCAB_ORIG} → {VOCAB_IMPR}", "label phrases in the rules, added by hand after "
                                                                  "reviewing the errors"),
                                  (f"{reg_o} → {reg_i}", "vet invoice receipt number: correct before the changes, "
                                                         "wrong after (a regression)"),
                                  ("New layouts", "may need more rule changes and retesting of every receipt")]):
    text(s, 5.3 + k * 2.6, 1.95, 2.45, 2.75, [dict(t=big, size=22, bold=True, color=RED, after=8),
                                             dict(t=small, size=13)], margin=0.2, fill=LIGHT)
text(s, 5.3, 4.85, 7.64, 1.1, [
    dict(t="AOAI used the same prompt for every receipt. Whether that means less maintenance on new layouts still "
           "needs to be tested.", size=14),
])

# ===================================================================== 6. where errors remain
ocr_read("receipt9", "FPatient")
s = content_slide(
    "Where Errors Remain: Both Approaches Make Mistakes",
    "On the vet invoice, the improved rules and AOAI each got a different field wrong.",
    "Receipt 9 (vet invoice). The improved rules now take the payment table's 'Receipt #' (380048) instead of the "
    "invoice number. Both OCR engines read 'FPatient:' (a stray mark merged into the label), so no rule matched the "
    "patient label; AOAI still found the name. AOAI chose the sub total (1361.50) instead of the total (1361.00). "
    "Values: detail.csv, checked against the extraction JSON. " + SOURCE)
top = make_crop("receipt9", ["Invoice No", "FClinician"], "crop_receipt9_top", ["Invoice No", "FClinician"])
pay = make_crop("receipt9", ["Sub Total", "HKD 1361.50", "Total Amount", "HKD 1361.00", "Details of Payments Received",
                             "Receipt #", "380048"], "crop_receipt9_pay",
                ["Sub Total", "HKD 1361.50", "Total Amount", "HKD 1361.00", "Receipt #", "380048"])
_, _, _, h1 = picture(s, top, 0.39, 1.95, 5.6, 1.0)
picture(s, pay, 0.39, 1.95 + h1 + 0.15, 5.6, 3.95 - h1 - 0.15)
example_table(s, 6.3, 1.95, "receipt9", ["receipt_number", "patient_name", "total_amount"],
              col_w=(1.15, 1.4, 1.35, 1.35, 1.39), row_h=0.68)
text(s, 6.3, 4.75, 6.64, 1.2, [
    dict(t="Improved rules: picked the payment receipt number.", size=13, bullet=True, after=4),
    dict(t="Rules: the OCR read \"FPatient:\", so the patient label was missed.", size=13, bullet=True, after=4),
    dict(t="AOAI: picked the sub total instead of the total.", size=13, bullet=True),
])

# ===================================================================== 7. practical comparison
s = content_slide(
    "Practical Comparison of the Two Candidates",
    "Similar accuracy on these receipts; they differ in cost, speed and how new layouts are handled.",
    f"Accuracy and time from summary.csv. AOAI tokens are averages over the {N_RECEIPTS} receipts. Prices for "
    "compute, Document Intelligence pages and gpt-5-mini tokens are not yet confirmed. " + SOURCE)
tbl = table(s, 0.39, 1.95, [3.1, 4.7, 4.75], [0.6] + [0.56] * 6)
header_row(tbl, ["", "PaddleOCR Light + improved rules", "Azure DI + AOAI"], size=15)
comp = [("Accuracy (these 9 receipts)", f"{acc(IMPR['light'])} ({count(IMPR['light'])}), tuned",
         f"{acc(AOAI)} ({count(AOAI)}), prompt unchanged"),
        ("Time per receipt", secs(IMPR["light"]), secs(AOAI)),
        ("Where it runs", "Our own compute, offline", "Azure cloud services"),
        ("Cost per receipt", "CPU time only, no API fee", f"DI page + about {TOKENS_IN} input / {TOKENS_OUT} output "
                                                          "tokens (prices to confirm)"),
        ("New receipt layouts", "Edit rules and retest", "Same prompt; still to be tested"),
        ("Main risk", "Maintenance and regressions", "API cost and model variability")]
for i, row in enumerate(comp, start=1):
    for c, v in enumerate(row):
        fill_cell(tbl.cell(i, c), [dict(t=v, size=14, bold=(c == 0))])
        cell_border(tbl.cell(i, c), sides="B")
note(s, "Both approaches read the receipt text well; the differences are in finding the fields.", y=6.1)

# ===================================================================== 8. conclusion
s = content_slide(
    "Conclusion and Next Steps",
    "Rules can reach high accuracy, but only with ongoing manual work. That is the decision to make.",
    f"Main finding: PaddleOCR Light reached {acc(IMPR['light'])} and Azure DI {acc(IMPR['di'])} only after the "
    f"rules were rewritten for these receipts; Azure DI + AOAI reached {acc(AOAI)} with no change. The benchmark "
    f"is small ({N_RECEIPTS} receipts) and the improved rules were tuned on it, so this is not a production "
    "decision. " + SOURCE)
text(s, 0.39, 1.95, 12.55, 1.55, [
    dict(t="The question is not only which method is most accurate.", size=19, bold=True, after=6),
    dict(t="It is whether we want to keep updating extraction rules each time an unfamiliar receipt layout "
           "appears, or pay per request for a model that may handle new layouts without changes.", size=16),
], margin=0.3, fill=LIGHT)
steps = [("1. Freeze both", "Lock the improved rules and the AOAI prompt."),
         ("2. Test on new receipts", "30+ real, anonymised receipts not seen before. Count the rule changes needed."),
         ("3. Confirm prices", "Compute, Document Intelligence pages and gpt-5-mini tokens.")]
sw = (12.55 - 2 * 0.3) / 3
for i, (head, desc) in enumerate(steps):
    text(s, 0.39 + i * (sw + 0.3), 3.75, sw, 1.4, [dict(t=head, size=16, bold=True, color=RED, after=6),
                                                   dict(t=desc, size=14)], margin=0.22, fill=LIGHT)
note(s, f"Limitation: {N_RECEIPTS} receipts, mostly printed mock receipts; the improved rules were tuned on them. "
        "Candidates for further testing: PaddleOCR Light + improved rules, and Azure DI + AOAI.", y=5.45)

prs.save(OUT)
print(f"Saved {OUT} ({len(prs.slides)} slides); vocabulary {VOCAB_ORIG} -> {VOCAB_IMPR}")
